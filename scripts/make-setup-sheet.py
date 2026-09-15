"""Regenerates Sharan's setup workbook (registers/setup/, git-ignored). Run: python scripts/make-setup-sheet.py (needs openpyxl).

Pre-filled from Sharan's WhatsApp centre message of 2026-09-15 (docs/academy/sharan-centre-details-2026-09-15.md).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, GOLD, CREAM, AMBER = "0D1F4B", "C9A227", "FFF8E1", "B45309"
OUT = Path(__file__).resolve().parents[1] / "registers" / "setup" / "Gurukul-Setup-Sheet.xlsx"

NK, PY, RD, BB, OPG = "NK Bagrodia Public School", "Play Yard", "R.D. Rajpal School", "Bal Bharati Public School", "OPG World School"
CENTRES = [
    # app name, sector, maps pin, days, note
    (NK, "Sector 4", "https://maps.app.goo.gl/V3n6Ln8ArezRh7BbA", "Mon, Wed, Fri", ""),
    (PY, "Sector 7", "https://maps.app.goo.gl/cNJwA7YPkBqxBx5H8", "Mon, Wed, Fri", 'Brochure calls it "Play Yard Sports" — rename in the app?'),
    (RD, "Sector 9", "https://maps.app.goo.gl/GJVvXDh2cX8c9ruB7", "Mornings Mon, Wed, Fri · Evenings Tue, Thu, Sat", ""),
    (BB, "Sector 12", "https://maps.app.goo.gl/uh8MJ1j3Yb2vjbMU7", "Tue, Thu, Sat", 'Brochure calls it "Bal Bharati School" — rename in the app?'),
    (OPG, "Sector 19B", "https://maps.app.goo.gl/HMeNprymcSUe9u5w8", "Mon to Fri", ""),
]
# Gurukul's own age groups (the app's list is aligned to these in Phase 2A).
AGES = ["U8", "U10", "U12", "U13", "U14", "U15", "U16", "U18", "U19", "SENIOR", "ELITE"]
MWF, TTS, WEEK = "Mon, Wed, Fri", "Tue, Thu, Sat", "Mon, Tue, Wed, Thu, Fri"
OPG_FEE = "₹3000 (5 days a week) / ₹2500 (3 days a week)"
BATCHES = [
    # centre, name, days, start, end, age, fee-if-different, please-confirm
    (NK, "Junior 5:30-6:30pm", MWF, "5:30 pm", "6:30 pm", "U12", "", "Message header says U8, U13, U19 but the details say U12, U16, Elite. Which is right?"),
    (NK, "Senior 6:30-7:30pm", MWF, "6:30 pm", "7:30 pm", "U16", "", ""),
    (NK, "Elite 7:30-8:30pm", MWF, "7:30 pm", "8:30 pm", "ELITE", "", "Advanced players. Same ₹2000 fee?"),
    (PY, "Junior 5:30-6:30pm", MWF, "5:30 pm", "6:30 pm", "U8", "", "Header says U8, U13; details say U8, U15."),
    (PY, "Senior 6:30-7:30pm", MWF, "6:30 pm", "7:30 pm", "U15", "", ""),
    (RD, "Morning 6:30-7:30am", MWF, "6:30 am", "7:30 am", "", "", "Which age group is the morning batch?"),
    (RD, "Junior 5:30-6:30pm", TTS, "5:30 pm", "6:30 pm", "U12", "", "Header says U8, U13, U19; details say U12, U19."),
    (RD, "Senior 6:30-7:30pm", TTS, "6:30 pm", "7:30 pm", "U19", "", ""),
    (BB, "Junior 5:30-6:30pm", TTS, "5:30 pm", "6:30 pm", "U12", "", "Header says U8, U13, U19; details say U12, U18."),
    (BB, "Senior 6:30-7:30pm", TTS, "6:30 pm", "7:30 pm", "U18", "", ""),
    (OPG, "Junior 5-6pm", WEEK, "5:00 pm", "6:00 pm", "U12", OPG_FEE,
     "The Sept register had 3 batches on Mon/Wed/Fri (4-5, 5-6, 6-7pm). This message has 2 batches Mon-Fri (5-6, 6-7pm). Which is current?"),
    (OPG, "Senior 6-7pm", WEEK, "6:00 pm", "7:00 pm", "U18", OPG_FEE, "Header says U8, U13, U18."),
]

thin = Side(style="thin", color="D0D5DD")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, intro, headers, widths, rows=(), blank=15, confirm_col=None):
    ws = wb.create_sheet(title)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = intro
    ws["A2"].font = Font(italic=True, color="555555")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(headers), 1))
    ws.row_dimensions[2].height = 48
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AMBER if i == confirm_col else NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = box
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[4].height = 34
    ws.freeze_panes = "B5"
    total = len(rows) + blank
    for r in range(total):
        vals = rows[r] if r < len(rows) else [""] * len(headers)
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=5 + r, column=ci, value=v)
            c.border = box
            c.alignment = wrap
            if ci == confirm_col and v:
                c.font = Font(bold=True, color=AMBER)
            elif r < len(rows):
                c.fill = PatternFill("solid", fgColor=CREAM)
    return ws, 5, 4 + total


def dropdown(ws, col, first, last, options, prompt):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.promptTitle, dv.prompt = "Pick one", prompt
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{last}")


wb = Workbook()

# 1. Start here
ws = wb.active
ws.title = "Start here"
lines = [
    ("Gurukul FC — Dashboard setup sheet", Font(bold=True, size=18, color=NAVY)),
    ("For Sharan. Takes about 15 minutes. Fill it on the phone in Google Sheets or on a laptop.", Font(italic=True, color="555555")),
    ("", None),
    ("How to fill", Font(bold=True, size=13, color=NAVY)),
    ("1. Go through the tabs at the bottom in order: Centres → Batches → Coaches → Fees → Past dues & advance.", None),
    ("2. Cream rows are already filled in from your WhatsApp centre message (15 Sep 2026). Correct anything that's wrong.", None),
    ("3. Orange text in the \"Please confirm\" column is a question for you. The message said two different things there.", None),
    ("4. Use the dropdowns where they appear. Leave a cell blank if you don't know.", None),
    ("5. When done, send the file back to Piyush. We'll load it into the app.", None),
    ("", None),
    ("Students are NOT filled in here", Font(bold=True, size=13, color=NAVY)),
    ("Once the batches are in the app, a Google Sheet with one tab per centre is created automatically.", None),
    ("Type each child there (name, batch, parent name, WhatsApp, date of birth). The app picks it up within 10 minutes.", None),
    ("", None),
    ("Privacy", Font(bold=True, size=13, color=NAVY)),
    ("This file will hold coach phone numbers. Share it only with Piyush, directly, not in group chats.", None),
]
for i, (text, font) in enumerate(lines, start=1):
    c = ws.cell(row=i, column=1, value=text)
    if font:
        c.font = font
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 110

# 2. Centres
sheet(
    wb, "Centres",
    "One row per ground, filled in from your message. Please check each Google Maps pin opens at the exact training spot. Coach check-in uses it later.",
    ["Centre name in the app", "Area / sector", "Google Maps pin link", "Training days", "Please confirm"],
    [28, 12, 44, 30, 44],
    rows=[[n, s, pin, days, note] for n, s, pin, days, note in CENTRES],
    blank=2,
    confirm_col=5,
)

# 3. Batches
ws, f, l = sheet(
    wb, "Batches",
    "One row per batch, filled in from your message. Add the head coach (and assistant, if any) for each batch, and answer the orange questions.",
    ["Centre", "Batch name", "Training days", "Start time", "End time", "Age group", "Head coach", "Assistant coach",
     "Monthly fee if different from the centre fee", "Please confirm"],
    [24, 20, 22, 10, 10, 10, 18, 18, 24, 46],
    rows=[[c, n, d, s, e, a, "", "", fee, q] for c, n, d, s, e, a, fee, q in BATCHES],
    blank=12,
    confirm_col=10,
)
dropdown(ws, "A", f, l, [n for n, *_ in CENTRES], "Which centre is this batch at?")
dropdown(ws, "F", f, l, AGES, "Age group of this batch")

# 4. Coaches
ws, f, l = sheet(
    wb, "Coaches",
    "Each coach logs in with their WhatsApp number and a temporary password that the app gives them. The Gmail address is optional and only needed for the \"Sign in with Google\" button.",
    ["Coach name", "WhatsApp number", "Role", "Gmail (optional)", "Centres / batches they coach", "Notes"],
    [22, 18, 16, 28, 34, 24],
    blank=12,
)
dropdown(ws, "C", f, l, ["Head coach", "Assistant coach"], "Head coaches can mark days off; assistants can't")

# 5. Fees
ws, f, l = sheet(
    wb, "Fees",
    "Fee rules. Filled in from your message where it gave an answer. Please fill the blank answers.",
    ["Question", "Answer", "Example / notes"],
    [52, 34, 50],
    rows=[
        [f"Monthly fee at {NK} (₹)", "2000", ""],
        [f"Monthly fee at {PY} (₹)", "2000", ""],
        [f"Monthly fee at {RD} (₹)", "2000", "Same for the morning batch?"],
        [f"Monthly fee at {BB} (₹)", "2000", ""],
        [f"Monthly fee at {OPG} (₹)", "3000 (5 days a week) / 2500 (3 days a week)", "Which OPG children come 3 days a week? Mark them in the student sheet later."],
        ["Does the fee differ by batch or age group?", "No", "Only OPG differs, by days per week"],
        ["Fee due on which day of the month?", "1", "1 = due on the 1st of every month"],
        ["Days of grace before a fee counts as overdue", "7", "Shown in red after this"],
        ["Sibling discount?", "", "e.g. 10% off for the 2nd child, or ₹500 off"],
        ["Any other standard discount?", "", "e.g. staff children, scholarship"],
        ["One-time registration / admission fee (₹)", "", "Leave blank if none"],
        ["Gurukul kit — T-shirt & shorts (₹, one-time)", "1000", "Compulsory for every new child?"],
        ["Cosco football (₹, optional)", "700", ""],
        ["Paytm / UPI ID shown in fee reminders", "", "e.g. gurukulfc@paytm"],
        ["Fee reminder language", "", "English, Hindi or both"],
        ["Can a child join mid-month at a reduced fee?", "", "Yes = charge only for the days left"],
    ],
    blank=0,
)
dropdown(ws, "B", f + 5, f + 5, ["Yes", "No"], "")
dropdown(ws, "B", f + 14, f + 14, ["English", "Hindi", "Both"], "")
dropdown(ws, "B", f + 15, f + 15, ["Yes", "No"], "")

# 6. Past dues & advance
ws, f, l = sheet(
    wb, "Past dues & advance",
    "Only list the exceptions: children who still owe money for months BEFORE September 2026, or who have paid ahead. Everyone else can be left out. (September payments are recorded in the app; OPG's are already in from the register.)",
    ["Student name", "Centre", "Batch", "Owes for which months?", "Amount owed (₹)", "Paid in advance up to (month)", "Notes"],
    [22, 24, 16, 22, 14, 22, 26],
    blank=25,
)
dropdown(ws, "B", f, l, [n for n, *_ in CENTRES], "")

for w in wb.worksheets:
    w.sheet_properties.tabColor = GOLD if w.title == "Start here" else NAVY

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("saved", OUT)
