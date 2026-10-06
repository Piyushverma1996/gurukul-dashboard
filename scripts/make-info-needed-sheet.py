"""Builds the "information still needed" workbook for Sharan (upload it to Google Sheets).

Run: python scripts/make-info-needed-sheet.py
Writes registers/setup/Gurukul-Information-Needed.xlsx (git-ignored: it lists real students).
Reads registers/roster/roster.csv for the students already loaded into the app.
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
ROSTER = ROOT / "registers" / "roster" / "roster.csv"
OUT = ROOT / "registers" / "setup" / "Gurukul-Information-Needed.xlsx"

NAVY, GOLD, CREAM, AMBER = "0E1B3D", "C9A227", "FFF8E1", "B45309"
thin = Side(style="thin", color="D0D5DD")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

CENTRE_BATCHES = {
    "NK Bagrodia Public School": ["Junior 5-6pm", "Senior 6-7pm", "Elite 7-8pm"],
    "Play Yard": ["Junior 5-6pm", "Senior 6-7pm"],
    "R.D. Rajpal School": ["Morning 6:30-7:30am", "Junior 5-6pm", "Senior 6-7pm"],
    "Bal Bharati Public School": ["Junior 5-6pm", "Senior 6-7pm"],
    "OPG World School": ["Junior 5-6pm", "Senior 6-7pm"],
}
STUDENT_COLS = ["Student name", "Parent name", "Parent WhatsApp", "Date of birth (DD/MM/YYYY)", "Gender",
                "Batch", "Monthly fee (₹)", "Joining date (DD/MM/YYYY)", "Anything we should know"]
STUDENT_WIDTHS = [24, 24, 20, 22, 10, 22, 14, 22, 30]


def sheet(wb, title, intro, headers, widths, rows=(), blank=15, highlight=None):
    ws = wb.create_sheet(title)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = intro
    ws["A2"].font = Font(italic=True, color="555555")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(headers), 1))
    ws.row_dimensions[2].height = 46
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AMBER if i == highlight else NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = box
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "B5"
    total = len(rows) + blank
    for r in range(total):
        vals = rows[r] if r < len(rows) else [""] * len(headers)
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=5 + r, column=ci, value=v)
            c.border = box
            c.alignment = wrap
            if r < len(rows) and ci != highlight:
                c.fill = PatternFill("solid", fgColor=CREAM)
            if ci == highlight and v:
                c.font = Font(bold=True, color=AMBER)
    return ws, 5, 4 + total


def dropdown(ws, col, first, last, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{last}")


wb = Workbook()
ws = wb.active
ws.title = "Start here"
for i, (text, font) in enumerate([
    ("Gurukul FC — information still needed", Font(bold=True, size=18, color=NAVY)),
    ("For Sharan. The app is live with Bal Bharati (40 students) and OPG (31 students). This file lists everything still missing.", Font(italic=True, color="555555")),
    ("", None),
    ("What to fill, in order of importance", Font(bold=True, size=13, color=NAVY)),
    ("1. Students at the three centres with no list yet: NK Bagrodia, Play Yard and RD Rajpal (one tab each).", None),
    ("2. Answers on the \"Questions\" tab — the Sector 9 morning batch, the UPI ID for fee reminders, and discounts.", None),
    ("3. \"Coaches\" tab: which coach runs which batch, so coaches see only their own children.", None),
    ("4. \"Check the 71 loaded\" tab: joining dates, and a quick check of what we read from your database.", None),
    ("5. \"Past dues & advance\" tab: only children who owe for earlier months or have paid ahead.", None),
    ("", None),
    ("How to use it", Font(bold=True, size=13, color=NAVY)),
    ("Upload this file to Google Drive and open it with Google Sheets. Fill it on your phone or laptop, then send it back to Piyush.", None),
    ("Leave anything you don't know blank. Cream cells are what the app already holds; the orange column is for your answer.", None),
    ("", None),
    ("Privacy", Font(bold=True, size=13, color=NAVY)),
    ("This file has children's and parents' details. Share it only with Piyush, directly, not in group chats.", None),
], start=1):
    c = ws.cell(row=i, column=1, value=text)
    if font:
        c.font = font
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 118

# --- Questions ---
qs = [
    ["Sector 9 morning batch (6:30–7:30 am): which age group?", "Loaded as U19 for now", "The app shows this age group on the batch and its students"],
    ["Sector 9 morning batch: same ₹2000 fee?", "Assumed yes", "Decides what parents are billed"],
    ["Elite batch (Sector 4, 7–8 pm): what is the monthly fee?", "Assumed ₹2000", "Decides what Elite parents are billed"],
    ["Which 3 days do OPG's 3-day students attend?", "", "So coaches only mark attendance on their days"],
    ["Paytm / UPI ID to show in fee reminder messages", "", "Goes into every WhatsApp fee reminder"],
    ["Which number should fee reminders come from?", "", "The WhatsApp number parents will see"],
    ["Fee reminder language: English, Hindi or both?", "", "Wording of the reminder message"],
    ["Sibling discount? (e.g. 10% off the second child)", "", "Applied automatically to the monthly fee"],
    ["Any other standard discount? (staff children, scholarships)", "", "Applied automatically"],
    ["One-time registration or admission fee (₹)?", "None assumed", "Not tracked yet — tell us if it exists"],
    ["Is the ₹1000 kit compulsory for every new child?", "", "We'll add one-time charges next; this decides who owes it"],
    ["Should the app track the ₹700 football purchases?", "", "Same as above"],
    ["Can a child join mid-month at a reduced fee?", "", "If yes, the app charges only the remaining days"],
    ["Rename \"Play Yard\" to \"Play Yard Sports\", and \"Bal Bharati Public School\" to \"Bal Bharati School\"?", "", "The names parents see in the app"],
    ["Which centres have Sharan's own WhatsApp group for parents?", "", "Used for rain / day-off messages later"],
]
sheet(wb, "Questions", "Answer in the orange column. Where the app already assumed something, it's shown in the middle column.",
      ["Question", "What the app assumes today", "Why it matters", "Your answer"], [58, 26, 40, 34],
      rows=[[q, a, w, ""] for q, a, w in qs], blank=4, highlight=4)

# --- Coaches ---
coach_rows = [["Deepanshu", "99589 76380", "Head coach", "", "", ""], ["Mom", "99119 22953", "Head coach", "", "", ""]]
ws, f, l = sheet(wb, "Coaches", "The app already has these two coaches. Add the rest, and tell us which batches each one runs — coaches only see their own batches.",
                 ["Coach name", "WhatsApp number", "Role", "Gmail (optional)", "Which centre(s) and batch(es) do they coach?", "Notes"],
                 [22, 18, 16, 26, 42, 24], rows=coach_rows, blank=12, highlight=5)
dropdown(ws, "C", f, l, ["Head coach", "Assistant coach", "Admin"])

# --- One student tab per centre with no roster ---
for centre in ["NK Bagrodia Public School", "Play Yard", "R.D. Rajpal School"]:
    short = {"NK Bagrodia Public School": "NK Bagrodia", "Play Yard": "Play Yard", "R.D. Rajpal School": "RD Rajpal"}[centre]
    ws, f, l = sheet(wb, f"Students – {short}",
                     f"Every child training at {centre}. The same columns as your database. Leave the fee blank for the standard ₹2000.",
                     STUDENT_COLS, STUDENT_WIDTHS, blank=45)
    dropdown(ws, "F", f, l, CENTRE_BATCHES[centre])
    dropdown(ws, "E", f, l, ["Male", "Female"])

# --- The 71 already loaded ---
rows = list(csv.DictReader(ROSTER.open(encoding="utf-8"))) if ROSTER.exists() else []
loaded = []
for r in rows:
    note = ""
    if r["name"] == "Girik Nadella":
        note = "Date of birth read 14/06/2107 in your file; we used 2017 — please confirm"
    if r["name"] == "Aarav Kumar":
        note = "Two children share this name; check both rows are right"
    dob = "/".join(reversed(r["dob"].split("-"))) if r["dob"] else ""
    loaded.append([r["name"], r["parent_name"], r["parent_phone"].replace("+91", ""), dob, r["gender"], r["centre"], r["batch"], r["monthly_fee"], "", note])
sheet(wb, "Check the 71 loaded",
      "These students are live in the app, taken from your database. Please add each child's joining date, and correct anything that's wrong. Everyone is currently recorded as joining on 1 October 2026.",
      ["Student name", "Parent name", "Parent WhatsApp", "Date of birth", "Gender", "Centre", "Batch", "Monthly fee (₹)",
       "Joining date (DD/MM/YYYY)", "Correction needed?"],
      [22, 22, 16, 14, 9, 24, 18, 13, 22, 40], rows=loaded, blank=5, highlight=9)

# --- Opening balances ---
ws, f, l = sheet(wb, "Past dues & advance",
                 "Only the exceptions: children who still owe money for months before October 2026, or who have already paid ahead. Everyone else can be left out.",
                 ["Student name", "Centre", "Owes for which months?", "Amount owed (₹)", "Paid in advance up to (month)", "Notes"],
                 [24, 26, 24, 16, 26, 28], blank=25)
dropdown(ws, "B", f, l, list(CENTRE_BATCHES))

for w in wb.worksheets:
    w.sheet_properties.tabColor = GOLD if w.title == "Start here" else NAVY
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"saved {OUT} ({len(loaded)} loaded students listed)")
